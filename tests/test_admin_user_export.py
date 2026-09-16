"""Student Accounts downloads: filter parity, spreadsheet safety, and access control."""

import csv
from datetime import datetime, timedelta, timezone
from io import BytesIO, StringIO
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import models
from app.admin_sessions import get_current_console_user
from app.database import Base, get_db
from app.routers import admin
from app.utils import user_export


@pytest.fixture
def export_context(monkeypatch):
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine, tables=[
        models.User.__table__, models.Subscription.__table__, models.AdminConsoleAccess.__table__,
    ])
    db = sessionmaker(bind=engine)()
    actor = models.User(id=9999, email="admin@example.com", is_admin=True, is_active=True, email_verified=True)
    app = FastAPI()
    app.include_router(admin.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_console_user] = lambda: actor
    monkeypatch.setattr(admin, "_is_admin_turnstile_required", lambda request: False)
    monkeypatch.setattr(admin, "check_ip_rate_limit", lambda **kwargs: (True, 0))
    with TestClient(app) as client:
        yield SimpleNamespace(db=db, app=app, actor=actor, client=client)
    db.close()
    engine.dispose()


def _add_user(db, index, **attributes):
    fields = dict(
        email=f"student{index}@example.com", username=f"student{index}",
        full_name=f"Student {index}", university="Example University",
        hashed_password="PASSWORD_MUST_NOT_BE_EXPORTED",
        created_at=datetime(2026, 9, 1, 12, 30, index % 60),
        email_notifications_enabled=False,
    )
    fields.update(attributes)
    user = models.User(**fields)
    db.add(user)
    db.commit()
    return user


def _csv_rows(response):
    assert response.status_code == 200, response.text
    return list(csv.DictReader(StringIO(response.content.decode("utf-8-sig"))))


@pytest.mark.parametrize("export_format", ["csv", "xlsx"])
def test_export_contains_all_rows_in_table_order(export_context, export_format):
    ctx = export_context
    for index in range(47):
        _add_user(ctx.db, index)
    listed = ctx.client.get("/api/admin/users").json()
    assert listed["total"] == 47
    assert len(listed["users"]) == 20

    response = ctx.client.get("/api/admin/users/export", params={"format": export_format})
    assert response.status_code == 200, response.text
    assert response.headers["cache-control"] == "no-store"
    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    assert response.headers["content-disposition"] == (
        f'attachment; filename="rilono-student-accounts-{stamp}.{export_format}"'
    )
    if export_format == "csv":
        assert response.content.startswith(b"\xef\xbb\xbf")
        assert response.headers["content-type"].startswith("text/csv")
        rows = _csv_rows(response)
        emails = [row["Email"] for row in rows]
        assert tuple(rows[0]) == user_export.HEADERS
        assert b"PASSWORD_MUST_NOT_BE_EXPORTED" not in response.content
    else:
        assert response.headers["content-type"] == user_export.XLSX_MEDIA_TYPE
        workbook = load_workbook(BytesIO(response.content))
        assert workbook.sheetnames == ["Student Accounts"]
        sheet = workbook.active
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == "A1:K48"
        assert tuple(cell.value for cell in sheet[1]) == user_export.HEADERS
        assert sheet["J2"].data_type == "d"
        assert isinstance(sheet["J2"].value, datetime)
        emails = [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)]
        assert all("PASSWORD_MUST_NOT_BE_EXPORTED" not in str(row) for row in sheet.values)
    assert len(emails) == 47
    assert emails == [f"student{index}@example.com" for index in reversed(range(47))]
    assert emails[:20] == [user["email"] for user in listed["users"]]


@pytest.mark.parametrize("filters,expected", [
    ({"plan": "visa_pass"}, ["student1@example.com", "student0@example.com"]),
    ({"plan": "free"}, ["student3@example.com", "student2@example.com"]),
    ({"role": "student"}, ["student2@example.com", "student0@example.com"]),
    ({"role": "staff"}, ["student3@example.com", "student1@example.com"]),
    ({"role": "developer"}, ["student3@example.com"]),
    ({"role": "admin"}, ["student1@example.com"]),
    ({"search": "needle", "plan": "visa_pass", "role": "student"}, ["student0@example.com"]),
    ({"search": "needle"}, ["student3@example.com", "student2@example.com", "student0@example.com"]),
])
def test_export_filters_match_list(export_context, filters, expected):
    ctx = export_context
    users = [
        _add_user(ctx.db, 0, full_name="Needle Student"),
        _add_user(ctx.db, 1, is_admin=True),
        _add_user(ctx.db, 2, university="Needle University"),
        _add_user(ctx.db, 3, is_developer=True, username="needle_developer"),
    ]
    ctx.db.add_all([
        models.Subscription(user_id=users[0].id, plan="pro", status="active"),
        models.Subscription(user_id=users[1].id, plan="pro", status="active"),
        models.Subscription(user_id=users[2].id, plan="pro", status="canceled"),
    ])
    ctx.db.commit()
    listed = ctx.client.get("/api/admin/users", params=filters).json()
    rows = _csv_rows(ctx.client.get("/api/admin/users/export", params={"format": "csv", **filters}))
    assert [row["Email"] for row in rows] == [user["email"] for user in listed["users"]] == expected


def test_unicode_csv_quoting_labels_and_name_fallback(export_context):
    ctx = export_context
    _add_user(
        ctx.db, 0, full_name='Zoë, "李"\nStudent', university="Université de Montréal",
        destination_country_code="CA", visa_type_key="ca_study_permit", acquisition_channel="google_organic",
        email_verified=True, is_active=False,
    )
    _add_user(ctx.db, 1, full_name=None)
    _add_user(ctx.db, 2, full_name=None, username=None)
    rows = _csv_rows(ctx.client.get("/api/admin/users/export?format=csv"))
    assert rows[0]["Name"] == "student2@example.com"
    assert rows[1]["Name"] == "student1"
    assert rows[2] == {
        "Name": 'Zoë, "李"\nStudent', "Email": "student0@example.com",
        "University": "Université de Montréal", "Acquisition source": "Google", "Country": "Canada",
        "Visa type": "Study Permit", "Role": "Student", "Status": "Inactive", "Verified": "Verified",
        "Created (UTC)": "2026-09-01 12:30:00", "Last login (UTC)": "",
    }


@pytest.mark.parametrize("value", ["=1+1", "+SUM(1,2)", "-1+2", "@SUM(1,2)", " \t=1+1", "\tplain", "\rplain", "\nplain"])
def test_user_text_cannot_execute_as_spreadsheet_formula(export_context, value):
    ctx = export_context
    _add_user(ctx.db, 0, full_name=value, university=value)
    rows = _csv_rows(ctx.client.get("/api/admin/users/export?format=csv"))
    assert rows[0]["Name"] == "'" + value
    assert rows[0]["University"] == "'" + value

    response = ctx.client.get("/api/admin/users/export?format=xlsx")
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    # XML serializers may normalize carriage returns into line feeds.
    assert sheet["A2"].value.replace("\r", "\n") == value.replace("\r", "\n")
    assert sheet["A2"].data_type == "s"
    assert sheet["C2"].value.replace("\r", "\n") == value.replace("\r", "\n")
    assert sheet["C2"].data_type == "s"
    assert all(cell.data_type != "f" for row in sheet for cell in row)


def test_excel_dates_are_normalized_to_utc_and_illegal_xml_is_removed():
    user = SimpleNamespace(
        full_name="Student\x00 Name", username=None, email="test@example.com", university=None,
        acquisition_channel=None, destination_country_code=None, visa_type_key=None,
        is_admin=False, is_developer=False, is_active=True, email_verified=False,
        created_at=datetime(2026, 9, 6, 17, 30, tzinfo=timezone(timedelta(hours=5, minutes=30))),
        last_login_at=None,
    )
    sheet = load_workbook(BytesIO(user_export.build_xlsx([user]))).active
    assert sheet["A2"].value == "Student Name"
    assert sheet["J2"].value == datetime(2026, 9, 6, 12)
    assert sheet["K2"].value is None


@pytest.mark.parametrize("export_format", ["csv", "xlsx"])
def test_empty_export_keeps_column_headers(export_context, export_format):
    response = export_context.client.get("/api/admin/users/export", params={"format": export_format})
    assert response.status_code == 200
    if export_format == "csv":
        rows = list(csv.reader(StringIO(response.content.decode("utf-8-sig"))))
    else:
        rows = list(load_workbook(BytesIO(response.content)).active.values)
    assert len(rows) == 1
    assert tuple(rows[0]) == user_export.HEADERS


@pytest.mark.parametrize("query,status_code", [
    ("format=pdf", 422), ("format=xls", 422), ("plan=unknown", 400),
    ("role=unknown", 400), ("search=", 422),
])
def test_invalid_export_filters_are_rejected(export_context, query, status_code):
    assert export_context.client.get(f"/api/admin/users/export?{query}").status_code == status_code


def test_export_requires_authenticated_admin(export_context):
    ctx = export_context
    ctx.actor.is_admin = False
    assert ctx.client.get("/api/admin/users/export").status_code == 403
    ctx.actor.is_developer = True
    assert ctx.client.get("/api/admin/users/export").status_code == 200
    ctx.app.dependency_overrides.pop(get_current_console_user)
    assert ctx.client.get("/api/admin/users/export").status_code == 401


def test_export_requires_valid_turnstile_proof(export_context, monkeypatch):
    ctx = export_context
    monkeypatch.setattr(admin, "_is_admin_turnstile_required", lambda request: True)
    assert ctx.client.get("/api/admin/users/export").status_code == 403
    ctx.client.cookies.set(admin.ADMIN_TURNSTILE_COOKIE_NAME, "invalid")
    assert ctx.client.get("/api/admin/users/export").status_code == 403
    ctx.client.cookies.set(admin.ADMIN_TURNSTILE_COOKIE_NAME, admin._build_turnstile_proof_token(ctx.actor.id + 1))
    assert ctx.client.get("/api/admin/users/export").status_code == 403
    ctx.client.cookies.set(admin.ADMIN_TURNSTILE_COOKIE_NAME, admin._build_turnstile_proof_token(ctx.actor.id))
    assert ctx.client.get("/api/admin/users/export").status_code == 200


def test_export_obeys_rate_limit(export_context, monkeypatch):
    calls = []

    def deny(**kwargs):
        calls.append(kwargs)
        return False, 23

    monkeypatch.setattr(admin, "check_ip_rate_limit", deny)
    response = export_context.client.get("/api/admin/users/export")
    assert response.status_code == 429
    assert response.headers["retry-after"] == "23"
    assert calls[0]["scope"] == "admin.users.export"
    assert calls[0]["extra_key"] == f"user:{export_context.actor.id}"


def test_excel_failure_returns_error_instead_of_mislabeled_csv(export_context, monkeypatch):
    def unavailable(users):
        raise ImportError("openpyxl")

    monkeypatch.setattr(user_export, "build_xlsx", unavailable)
    response = export_context.client.get("/api/admin/users/export?format=xlsx")
    assert response.status_code == 503
    assert "CSV" in response.json()["detail"]
    assert "content-disposition" not in response.headers
