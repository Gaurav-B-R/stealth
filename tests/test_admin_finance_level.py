"""Level 50/50: founder-to-founder settlements move founder spend and nothing else."""

from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import models
from app.admin_sessions import get_current_console_user
from app.database import Base, get_db
from app.routers import admin


@pytest.fixture
def ctx(monkeypatch):
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine, tables=[
        models.User.__table__, models.AdminConsoleAccess.__table__,
        models.CompanyFinanceEntry.__table__, models.SubscriptionPayment.__table__,
    ])
    db = sessionmaker(bind=engine)()
    actor = models.User(id=9999, email="admin@example.com", is_admin=True, is_active=True, email_verified=True)
    app = FastAPI()
    app.include_router(admin.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_console_user] = lambda: actor
    monkeypatch.setattr(admin, "_is_admin_turnstile_required", lambda request: False)
    monkeypatch.setattr(admin, "check_ip_rate_limit", lambda **kwargs: (True, 0))
    monkeypatch.setattr(admin.fx, "get_state", lambda force=False: {"rate": 84.0, "source": "fallback"})
    with TestClient(app) as client:
        yield SimpleNamespace(db=db, client=client)
    db.close()
    engine.dispose()


def _expense(db, amount, paid_by, *, vendor="Claude", occurred_on=date(2026, 9, 1)):
    row = models.CompanyFinanceEntry(
        entry_type="expense", category="AI Model Tools", vendor=vendor,
        amount_usd=-abs(Decimal(str(amount))), occurred_on=occurred_on, paid_by=paid_by, source="manual",
    )
    db.add(row)
    db.commit()
    return row


def _seed_screenshot_split(db):
    """The live split when this was built: Gaurav $871.33 vs Kushal $334.24."""
    _expense(db, "800.00", "Gaurav", vendor="Claude", occurred_on=date(2026, 8, 3))
    _expense(db, "71.33", "Gaurav", vendor="Gemini API", occurred_on=date(2026, 9, 18))
    _expense(db, "334.24", "Kushal", vendor="Render", occurred_on=date(2026, 7, 9))


def _analytics(ctx):
    response = ctx.client.get("/api/admin/company-finance/analytics")
    assert response.status_code == 200, response.text
    return response.json()


def _split(payload):
    return {item["label"]: item["amount_usd"] for item in payload["contributor_breakdown"]}


def _level(ctx, balance, **overrides):
    body = {
        "expected_payer": balance["payer"],
        "expected_transfer_usd": balance["transfer_usd"],
        "occurred_on": datetime.now(timezone.utc).date().isoformat(),
    }
    body.update(overrides)
    return ctx.client.post("/api/admin/company-finance/level", json=body)


def test_level_records_one_settlement_and_moves_only_the_founder_split(ctx):
    _seed_screenshot_split(ctx.db)
    before = _analytics(ctx)
    balance = before["founder_balance"]
    assert balance["is_level"] is False
    assert (balance["payer"], balance["payee"], balance["transfer_usd"]) == ("Kushal", "Gaurav", 268.55)
    assert balance["total_usd"] == 1205.57
    assert {f["label"]: f["after_usd"] for f in balance["founders"]} == {"Gaurav": 602.78, "Kushal": 602.79}

    today = datetime.now(timezone.utc).date().isoformat()
    response = _level(ctx, balance, occurred_on=today)
    assert response.status_code == 200, response.text
    entry = response.json()["entry"]
    assert entry["kind"] == "Settlement"
    assert (entry["paid_by"], entry["vendor"], entry["amount_usd"]) == ("Kushal", "Gaurav", 268.55)
    assert response.json()["founder_balance"]["is_level"] is True

    after = _analytics(ctx)
    assert _split(after) == {"Kushal": 602.79, "Gaurav": 602.78}
    assert after["founder_balance"]["is_level"] is True
    assert after["founder_balance"]["transfer_usd"] == 0
    # Not company money: totals, ROI, the monthly chart and the vendor breakdown are untouched.
    for key in ("total_invested_usd", "total_returns_usd", "net_usd", "roi_percent",
                "break_even_gap_usd", "investment_entry_count", "return_entry_count"):
        assert after["summary"][key] == before["summary"][key], key
    assert after["summary"]["founder_settlement_count"] == 1
    assert after["monthly_series"] == before["monthly_series"]
    assert after["expense_breakdown"] == before["expense_breakdown"]
    assert sum(_split(after).values()) == pytest.approx(before["summary"]["total_invested_usd"])
    assert [round(item["percentage"], 2) for item in after["contributor_breakdown"]] == [50.0, 50.0]

    settlement_rows = [row for row in after["ledger"] if row["kind"] == "Settlement"]
    assert len(settlement_rows) == 1
    row = settlement_rows[0]
    assert row["occurred_on"] == today
    assert (row["paid_by"], row["paid_to"], row["vendor"]) == ("Kushal", "Gaurav", "Gaurav")
    assert row["amount_usd"] == 268.55
    assert row["category"] == "Founder Settlement"
    assert row["source"] == "level"
    assert "Kushal paid Gaurav $268.55" in row["description"]
    # Newest first: dated today, it sits at the top of the ledger.
    assert after["ledger"][0]["id"] == row["id"]


def test_level_is_a_no_op_once_level_and_never_flips_the_odd_cent(ctx):
    _seed_screenshot_split(ctx.db)
    balance = _analytics(ctx)["founder_balance"]
    assert _level(ctx, balance).status_code == 200

    # A double-click resends the same (now stale) confirmation.
    again = _level(ctx, balance)
    assert again.status_code == 409
    assert "already 50/50" in again.json()["detail"]
    payload = _analytics(ctx)
    assert payload["founder_balance"]["is_level"] is True
    assert payload["summary"]["founder_settlement_count"] == 1


def test_stale_confirmation_is_refused_and_records_nothing(ctx):
    _seed_screenshot_split(ctx.db)
    balance = _analytics(ctx)["founder_balance"]
    # Someone adds an expense after the confirmation was shown.
    _expense(ctx.db, "10.00", "Gaurav")

    response = _level(ctx, balance)
    assert response.status_code == 409
    assert "changed since this page loaded" in response.json()["detail"]
    assert ctx.db.query(models.CompanyFinanceEntry).filter_by(entry_type="settlement").count() == 0

    wrong_payer = _level(ctx, _analytics(ctx)["founder_balance"], expected_payer="Gaurav")
    assert wrong_payer.status_code == 409
    assert ctx.db.query(models.CompanyFinanceEntry).filter_by(entry_type="settlement").count() == 0


def test_direction_follows_whoever_spent_less(ctx):
    _expense(ctx.db, "100.00", "Gaurav")
    _expense(ctx.db, "300.00", "Kushal")
    balance = _analytics(ctx)["founder_balance"]
    assert (balance["payer"], balance["payee"], balance["transfer_usd"]) == ("Gaurav", "Kushal", 100.0)
    assert _level(ctx, balance).status_code == 200
    assert _split(_analytics(ctx)) == {"Gaurav": 200.0, "Kushal": 200.0}


def test_new_spend_after_leveling_can_be_leveled_again(ctx):
    _expense(ctx.db, "100.00", "Gaurav")
    _expense(ctx.db, "50.00", "Kushal")
    assert _level(ctx, _analytics(ctx)["founder_balance"]).status_code == 200
    _expense(ctx.db, "40.00", "Kushal")
    balance = _analytics(ctx)["founder_balance"]
    assert (balance["payer"], balance["payee"], balance["transfer_usd"]) == ("Gaurav", "Kushal", 20.0)
    assert _level(ctx, balance).status_code == 200
    assert _split(_analytics(ctx)) == {"Gaurav": 95.0, "Kushal": 95.0}


def test_founder_names_match_case_insensitively(ctx):
    _expense(ctx.db, "90.00", "Gaurav")
    _expense(ctx.db, "10.00", " gaurav ")
    _expense(ctx.db, "20.00", "KUSHAL")
    payload = _analytics(ctx)
    assert _split(payload) == {"Gaurav": 100.0, "Kushal": 20.0}
    assert payload["founder_balance"]["transfer_usd"] == 40.0


def test_nothing_to_level_without_spend(ctx):
    balance = _analytics(ctx)["founder_balance"]
    assert balance["is_level"] is True
    response = _level(ctx, balance, expected_payer="Kushal", expected_transfer_usd=0)
    assert response.status_code == 409


@pytest.mark.parametrize("offset_days,status_code", [(-1, 200), (1, 200), (-2, 400), (30, 400)])
def test_settlement_date_must_be_the_click_day(ctx, offset_days, status_code):
    _expense(ctx.db, "100.00", "Gaurav")
    balance = _analytics(ctx)["founder_balance"]
    when = datetime.now(timezone.utc).date() + timedelta(days=offset_days)
    response = _level(ctx, balance, occurred_on=when.isoformat())
    assert response.status_code == status_code, response.text
    if status_code == 200:
        assert response.json()["entry"]["occurred_on"] == when.isoformat()


def test_omitted_date_defaults_to_utc_today(ctx):
    _expense(ctx.db, "100.00", "Gaurav")
    balance = _analytics(ctx)["founder_balance"]
    response = _level(ctx, balance, occurred_on=None)
    assert response.status_code == 200, response.text
    assert response.json()["entry"]["occurred_on"] == datetime.now(timezone.utc).date().isoformat()


def test_settlements_cannot_be_forged_or_edited_but_delete_undoes_them(ctx):
    _seed_screenshot_split(ctx.db)
    forged = ctx.client.post("/api/admin/company-finance/entries", json={
        "occurred_on": "2026-10-01", "entry_type": "settlement", "category": "Founder Settlement",
        "vendor": "Gaurav", "amount_usd": 50, "paid_by": "Kushal",
    })
    assert forged.status_code == 400

    entry_id = _level(ctx, _analytics(ctx)["founder_balance"]).json()["entry"]["id"]
    edited = ctx.client.patch(f"/api/admin/company-finance/entries/{entry_id}", json={"amount_usd": 1})
    assert edited.status_code == 400
    assert "can't be edited" in edited.json()["detail"]

    expense_id = ctx.db.query(models.CompanyFinanceEntry).filter_by(entry_type="expense").first().id
    converted = ctx.client.patch(f"/api/admin/company-finance/entries/{expense_id}", json={"entry_type": "settlement"})
    assert converted.status_code == 400

    assert ctx.client.delete(f"/api/admin/company-finance/entries/{entry_id}").status_code == 200
    payload = _analytics(ctx)
    assert _split(payload) == {"Gaurav": 871.33, "Kushal": 334.24}
    assert payload["founder_balance"]["transfer_usd"] == 268.55


def test_export_lists_the_settlement_and_the_level_founder_split(ctx):
    _seed_screenshot_split(ctx.db)
    assert _level(ctx, _analytics(ctx)["founder_balance"]).status_code == 200
    response = ctx.client.get("/api/admin/company-finance/export", params={"format": "xlsx"})
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    ledger = list(workbook["Ledger"].iter_rows(min_row=2, values_only=True))
    settlements = [row for row in ledger if row[1] == "Settlement"]
    assert len(settlements) == 1
    assert settlements[0][2:5] == ("Gaurav", "Founder Settlement", "Kushal")
    # Not company money: 0 in the signed Amount column, the transfer in its own column.
    assert settlements[0][6] == 0
    assert settlements[0][9] in (None, "")
    assert settlements[0][10] == 268.55
    net = ctx.client.get("/api/admin/company-finance/analytics").json()["summary"]["net_usd"]
    assert round(sum(row[6] for row in ledger), 2) == net == -1205.57
    by_founder = {row[0]: row[1] for row in workbook["Spend by founder"].iter_rows(min_row=2, values_only=True)}
    assert by_founder == {"Kushal": 602.79, "Gaurav": 602.78}
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(min_row=5, max_col=2, values_only=True) if row[0]}
    assert summary["Total invested (USD)"] == 1205.57


def test_csv_export_amount_column_still_sums_to_net(ctx):
    _seed_screenshot_split(ctx.db)
    assert _level(ctx, _analytics(ctx)["founder_balance"]).status_code == 200
    response = ctx.client.get("/api/admin/company-finance/export", params={"format": "csv"})
    assert response.status_code == 200, response.text
    lines = response.content.decode("utf-8-sig").splitlines()
    header_index = next(i for i, line in enumerate(lines) if line.startswith("Date (UTC)"))
    import csv
    rows = list(csv.DictReader(lines[header_index:]))
    assert round(sum(float(row["Amount (USD)"]) for row in rows), 2) == -1205.57
    settlement = next(row for row in rows if row["Type"] == "Settlement")
    assert settlement["Founder transfer (USD)"] == "268.55"


@pytest.mark.parametrize("label", ["Both", "Gaurav B R", "Unassigned-ish"])
def test_spend_under_other_names_blocks_level_until_reassigned(ctx, label):
    _seed_screenshot_split(ctx.db)
    stray = _expense(ctx.db, "100.00", label)
    payload = _analytics(ctx)
    balance = payload["founder_balance"]
    assert balance["other_contributors"] == [{"label": label, "amount_usd": 100.0, "percentage": 0}]
    assert payload["summary"]["total_invested_usd"] == 1305.57

    response = _level(ctx, balance)
    assert response.status_code == 409
    assert f"{label} $100.00" in response.json()["detail"]
    assert ctx.db.query(models.CompanyFinanceEntry).filter_by(entry_type="settlement").count() == 0

    # Re-attributing the entry to a founder unblocks it.
    fixed = ctx.client.patch(f"/api/admin/company-finance/entries/{stray.id}", json={"paid_by": "Gaurav"})
    assert fixed.status_code == 200, fixed.text
    balance = _analytics(ctx)["founder_balance"]
    assert balance["other_contributors"] == []
    assert (balance["payer"], balance["transfer_usd"]) == ("Kushal", 318.55)
    assert _level(ctx, balance).status_code == 200
    assert _split(_analytics(ctx)) == {"Gaurav": 652.78, "Kushal": 652.79}


def test_negative_founder_spend_after_a_deleted_expense_is_reported_not_hidden(ctx):
    _seed_screenshot_split(ctx.db)
    assert _level(ctx, _analytics(ctx)["founder_balance"]).status_code == 200
    big = ctx.db.query(models.CompanyFinanceEntry).filter_by(vendor="Claude").first()
    assert ctx.client.delete(f"/api/admin/company-finance/entries/{big.id}").status_code == 200
    payload = _analytics(ctx)
    shares = {item["label"]: item["percentage"] for item in payload["contributor_breakdown"]}
    assert shares == {"Kushal": 148.63, "Gaurav": -48.63}
    balance = payload["founder_balance"]
    assert (balance["payer"], balance["payee"], balance["transfer_usd"]) == ("Gaurav", "Kushal", 400.01)
