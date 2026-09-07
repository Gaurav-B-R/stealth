"""Safe spreadsheet serialization of the Student Accounts table's public columns."""

import csv
import re
from datetime import datetime, timezone
from io import BytesIO, StringIO

from app import acquisition


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS = (
    "Name", "Email", "University", "Acquisition source", "Country", "Visa type",
    "Role", "Status", "Verified", "Created (UTC)", "Last login (UTC)",
)

# Match the concise labels displayed by static/admin_console.js.
COUNTRY_LABELS = {"US": "United States", "UK": "United Kingdom", "CA": "Canada", "AU": "Australia"}
VISA_LABELS = {
    "us_f1": "F-1 Student", "us_j1": "J-1 Exchange", "us_m1": "M-1 Vocational",
    "uk_student": "Student Visa", "uk_short_study": "Short-Term Study",
    "ca_study_permit": "Study Permit", "ca_sds": "SDS", "au_subclass500": "Subclass 500",
}
SOURCE_LABEL_OVERRIDES = {
    "google_organic": "Google", "gemini": "AI Referral", "claude": "AI Referral",
    "email": "Email", "referral": "Referral", "direct": "Direct", "other": "Other",
    "untracked": "",
}
# XML 1.0 forbids these control characters. Preserve legitimate tabs/newlines.
INVALID_XML_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def _utc_datetime(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    # Database timestamps without an offset already represent UTC. Excel stores
    # timezone-free date cells, so normalize aware values before removing the offset.
    return value.astimezone(timezone.utc).replace(tzinfo=None) if value.tzinfo else value


def _rows(users):
    for user in users:
        country_code = str(user.destination_country_code or "").upper()
        country = COUNTRY_LABELS.get(country_code, "")
        channel = user.acquisition_channel
        source = SOURCE_LABEL_OVERRIDES.get(
            channel, acquisition.CHANNEL_META.get(channel, {}).get("label", "Other")
        ) if channel else ""
        values = (
            user.full_name or user.username or user.email or "Unknown user",
            user.email or "",
            user.university or "",
            source,
            country,
            VISA_LABELS.get(user.visa_type_key, "") if country else "",
            "Developer" if user.is_developer else "Admin" if user.is_admin else "Student",
            "Active" if user.is_active else "Inactive",
            "Verified" if user.email_verified else "Pending",
            _utc_datetime(user.created_at),
            _utc_datetime(user.last_login_at),
        )
        yield [INVALID_XML_CHARACTERS.sub("", value) if isinstance(value, str) else value for value in values]


def _csv_value(value):
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, str) and (
        value.startswith(("\t", "\r", "\n")) or value.lstrip().startswith(("=", "+", "-", "@"))
    ):
        # CSV has no cell types: prefix potentially executable input with an apostrophe
        # so spreadsheet apps treat it as text, including whitespace-prefixed formulas.
        return "'" + value
    return value


def build_csv(users) -> bytes:
    buffer = StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(HEADERS)
    writer.writerows([_csv_value(value) for value in row] for row in _rows(users))
    return buffer.getvalue().encode("utf-8-sig")


def build_xlsx(users) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Student Accounts")
    sheet.freeze_panes = "A2"
    for index, width in enumerate((28, 36, 30, 23, 23, 23, 15, 13, 13, 23, 23), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    headers = []
    for title in HEADERS:
        cell = WriteOnlyCell(sheet, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4338CA")
        headers.append(cell)
    sheet.append(headers)
    row_count = 1
    for values in _rows(users):
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                # Explicit text prevents Excel formulas without altering the actual data.
                cell.data_type = "s"
            elif isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            cells.append(cell)
        sheet.append(cells)
        row_count += 1
    sheet.auto_filter.ref = f"A1:K{row_count}"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
